import csv
from pathlib import Path
from typing import Any, Dict, List, Optional

import openpyxl


def read_csv_file(file_path: str) -> List[Dict[str, Any]]:
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"Файл не найден: {file_path}")

    transactions: List[Dict[str, Any]] = []

    with path.open("r", encoding="utf-8", newline="") as file:
        reader = csv.DictReader(file, delimiter=";")

        if reader.fieldnames is None:
            raise ValueError(f"CSV-файл пуст или не содержит заголовков: {file_path}")

        if all(field is None or field.strip() == "" for field in reader.fieldnames):
            raise ValueError(f"Заголовки CSV-файла пустые: {file_path}")

        for row in reader:
            cleaned_row: Dict[str, Any] = {k: v for k, v in row.items()}
            transactions.append(cleaned_row)

    if not transactions:
        raise ValueError(f"CSV-файл не содержит данных: {file_path}")

    return transactions


def read_xlsx_file(file_path: str) -> List[Dict[str, Any]]:
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"Файл не найден: {file_path}")

    workbook: Optional[openpyxl.Workbook] = None
    try:
        workbook = openpyxl.load_workbook(path, data_only=True)
        sheet = workbook.active

        if sheet is None:
            raise ValueError(f"В XLSX-файле отсутствует активный лист: {file_path}")

        if sheet.max_row < 2:
            raise ValueError(f"XLSX-файл не содержит данных (только заголовки или пуст): {file_path}")

        headers: List[str] = []
        first_row = sheet[1]
        for cell in first_row:
            header_value = cell.value
            if header_value is None:
                header_value = ""
            headers.append(str(header_value))

        if all(h == "" for h in headers):
            raise ValueError(f"Заголовки XLSX-файла пустые: {file_path}")

        transactions: List[Dict[str, Any]] = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if all(value is None for value in row):
                continue

            transaction: Dict[str, Any] = {}
            for i, header in enumerate(headers):
                if i < len(row):
                    transaction[header] = row[i]
                else:
                    transaction[header] = None

            transactions.append(transaction)

        if not transactions:
            raise ValueError(f"XLSX-файл не содержит данных после пропуска пустых строк: {file_path}")

        return transactions

    except FileNotFoundError:
        raise
    except Exception as e:
        raise ValueError(f"Ошибка при чтении XLSX-файла: {e}") from e
    finally:
        if workbook is not None:
            workbook.close()
