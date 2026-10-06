import csv
import json
import logging
from pathlib import Path
from typing import Any, Dict, List

from openpyxl import load_workbook

MODULE_NAME = "readers"

logger = logging.getLogger(MODULE_NAME)
if not logger.handlers:
    logger.setLevel(logging.DEBUG)
    file_handler = logging.FileHandler("readers.log", encoding="utf-8")
    file_handler.setLevel(logging.DEBUG)
    file_formatter = logging.Formatter(
        "%(asctime)s - %(name)s - %(levelname)s - %(message)s",
        datefmt="%Y-%m-%d %H:%M:%S",
    )
    file_handler.setFormatter(file_formatter)
    logger.addHandler(file_handler)


def read_json_file(path: str) -> List[Dict[str, Any]]:
    try:
        with open(path, encoding="utf-8") as f:
            data = json.load(f)
        if isinstance(data, list):
            logger.info("Successfully read JSON file: %s, rows=%d", path, len(data))
            return data
        return []
    except (FileNotFoundError, json.JSONDecodeError) as e:
        logger.error("Error reading JSON file %s: %s", path, e)
        return []


def read_csv(path: str) -> List[Dict[str, Any]]:
    file_path = Path(path)
    if not file_path.exists():
        msg = f"CSV file not found: {path}"
        logger.error(msg, exc_info=True)
        raise FileNotFoundError(msg)

    with open(path, encoding="utf-8") as f:
        reader = csv.DictReader(f, delimiter=";")
        result: List[Dict[str, Any]] = list(reader)

    logger.info("Successfully read CSV file: %s, rows=%d", path, len(result))
    return result


def read_excel(path: str) -> List[Dict[str, Any]]:
    file_path = Path(path)
    if not file_path.exists():
        msg = f"Excel file not found: {path}"
        logger.error(msg, exc_info=True)
        raise FileNotFoundError(msg)

    wb = load_workbook(path)
    ws = wb.active

    headers = [cell.value for cell in ws[1]]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    result: List[Dict[str, Any]] = [dict(zip(headers, row)) for row in rows]

    logger.info("Successfully read Excel file: %s, rows=%d", path, len(result))
    return result


read_csv_transactions = read_csv
read_excel_transactions = read_excel

__all__ = [
    "read_json_file",
    "read_csv",
    "read_excel",
    "read_csv_transactions",
    "read_excel_transactions",
]
