import logging
from pathlib import Path
from typing import Any, Dict, List

from openpyxl import load_workbook

MODULE_NAME = "readers.excel"
logger = logging.getLogger(MODULE_NAME)

if not logger.handlers:
    logger.setLevel(logging.DEBUG)
    file_handler = logging.FileHandler("readers.log", encoding="utf-8")
    file_formatter = logging.Formatter(
        "%(asctime)s - %(name)s - %(levelname)s - %(message)s",
        datefmt="%Y-%m-%d %H:%M:%S",
    )
    file_handler.setFormatter(file_formatter)
    logger.addHandler(file_handler)


def read_excel(path: str) -> List[Dict[str, Any]]:
    file_path = Path(path)
    if not file_path.exists():
        msg = f"Excel file not found: {path}"
        logger.error(msg, exc_info=True)
        raise FileNotFoundError(msg)

    wb = load_workbook(path)
    ws = wb.active

    headers = [cell.value for cell in ws[1]]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    result: List[Dict[str, Any]] = [dict(zip(headers, row)) for row in rows]

    logger.info("Successfully read Excel file: %s, rows=%d", path, len(result))
    return result
